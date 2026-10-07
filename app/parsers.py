"""Parsers for every file type in the corpus.

Contract: parse_file() either returns a list of chunks or raises. The caller
(indexer.build_index) catches per file, so one hostile file never stops the walk.

A chunk is {"text": str, "loc": str}. The indexer adds file / kind / status.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

csv.field_size_limit(10_000_000)

MAX_CHARS = 1400

TEXT_EXT = {".txt", ".log", ".py", ".md", ".rst", ".json", ".yaml", ".yml",
            ".xml", ".ini", ".cfg", ".toml"}
TABLE_EXT = {".csv", ".tsv"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}

IMAGE_PROMPT = (
    "Transcribe ALL text visible in this image exactly as printed, line by line. "
    "Keep structure: for tables and diagrams write each row/pin/label together with "
    "its neighbouring values (e.g. 'A7 = SIGNAL_NAME'). Then add one sentence "
    "describing what the image shows."
)


# --------------------------------------------------------------------------- utils

def chunk_lines(lines: list[str], max_chars: int = MAX_CHARS, overlap: int = 3):
    """Yield (first_line_no, last_line_no, text) windows of at most ~max_chars."""
    """
    	Splits text into windows of about 1,400 characters. A window starts 3 lines before the previous one ended, so a fact on a boundary isn't cut off.
        Lines longer than 1,400 characters are split first. It returns (first_line, last_line, text).
    """
    flat: list[str] = []
    for ln in lines:
        while len(ln) > max_chars:
            flat.append(ln[:max_chars])
            ln = ln[max_chars:]
        flat.append(ln)
    out, i, n = [], 0, len(flat)
    while i < n:
        j, size = i, 0
        while j < n and (j == i or size + len(flat[j]) + 1 <= max_chars):
            size += len(flat[j]) + 1
            j += 1
        text = "\n".join(flat[i:j]).strip()
        if text:
            out.append((i + 1, j, text))
        if j >= n:
            break
        i = max(j - overlap, i + 1)
    return out


def read_text(path: Path) -> str:
    """
    Reads a file as text. If the first 8 KB contains a NUL byte, the file is binary and it raises an error.
    Otherwise it decodes as UTF-8, replacing bad characters.
    """
    with open(path, "rb") as f:
        data = f.read(50_000_000)
    if b"\x00" in data[:8192]:
        raise ValueError("binary file")
    return data.decode("utf-8-sig", errors="replace")


def _cell(v) -> str:
    """
    Turns a spreadsheet value into a clean string. None becomes "",
    4.0 becomes 4, dates become ISO text, and extra whitespace is collapsed.
    """
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return re.sub(r"\s+", " ", str(v)).strip()


def table_chunks(sheet: str, rows) -> list[dict]:
    """Rows -> chunks of 'column: value' lines, each chunk repeating the header."""
    """
    Shared by CSV and XLSX. The first row is the header if it has at least 2 filled cells, otherwise columns are named col1, col2... Each row becomes [sheet] row 5: Part number: X; Product: TQ-40. Rows are grouped, at most 8 per chunk,
    and every chunk repeats the header line so the model knows what each column means.
    """
    rows = [[_cell(c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    header = rows[0]
    body = rows[1:]
    if sum(1 for c in header if c) < 2 or not body:  # no usable header row
        header, body = [], rows
    width = max(len(r) for r in rows)
    head = [(header[i] if i < len(header) and header[i] else f"col{i + 1}") for i in range(width)]

    lines = []
    for n, r in enumerate(body, 2 if header else 1):
        cells = "; ".join(f"{head[i]}: {v}" for i, v in enumerate(r) if v)
        lines.append((n, f"[{sheet}] row {n}: {cells}"))

    prefix = f"[{sheet}] columns: " + " | ".join(head)
    out, cur, size = [], [], 0
    for n, line in lines:
        if cur and (size + len(line) > MAX_CHARS or len(cur) >= 8):
            out.append({"text": prefix + "\n" + "\n".join(l for _, l in cur),
                        "loc": f"{sheet} rows {cur[0][0]}-{cur[-1][0]}"})
            cur, size = [], 0
        cur.append((n, line))
        size += len(line)
    if cur:
        out.append({"text": prefix + "\n" + "\n".join(l for _, l in cur),
                    "loc": f"{sheet} rows {cur[0][0]}-{cur[-1][0]}"})
    return out


# --------------------------------------------------------------------------- formats

def parse_text(path: Path) -> list[dict]:
    """
    For .txt, .log, .py and similar. It reads the file and calls chunk_lines. The location is lines a-b.
    """
    text = read_text(path)
    return [{"text": t, "loc": f"lines {a}-{b}"} for a, b, t in chunk_lines(text.splitlines())]


def parse_csv(path: Path) -> list[dict]:
    """
    Reads the text and detects the delimiter (tab for .tsv, otherwise it sniffs `, ; tab
    """
    text = read_text(path)
    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    if delim == ",":
        try:
            delim = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            pass
    return table_chunks(path.stem, csv.reader(io.StringIO(text), delimiter=delim))


def parse_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    """
    Opens with openpyxl in read-only mode, using cached formula values. It reads every sheet, including hidden ones, with table_chunks.
    """
    wb = load_workbook(str(path), read_only=True, data_only=True)  # every sheet, hidden too
    out = []
    try:
        for ws in wb.worksheets:
            out.extend(table_chunks(ws.title, ws.iter_rows(values_only=True)))
    finally:
        wb.close()
    return out


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_walk(el, out: list[str]) -> None:
    """
    	Recursive helper for Word XML. A paragraph becomes its text,
        a table row becomes cell | cell | cell, and any other element is searched inside.
    """
    for ch in el:
        if ch.tag == _W + "p":
            parts = []
            for n in ch.iter():
                if n.tag == _W + "t" and n.text:
                    parts.append(n.text)
                elif n.tag in (_W + "tab", _W + "br"):
                    parts.append(" ")
            t = "".join(parts).strip()
            if t:
                out.append(t)
        elif ch.tag == _W + "tbl":
            for tr in ch.findall(_W + "tr"):
                cells = []
                for tc in tr.findall(_W + "tc"):
                    sub: list[str] = []
                    _docx_walk(tc, sub)
                    cells.append(" ".join(sub))
                if any(cells):
                    out.append(" | ".join(cells))
        else:
            _docx_walk(ch, out)


def parse_docx(path: Path) -> list[dict]:
    """
    	Opens the .docx as a zip and reads the body, headers, footers and footnotes.
        An encrypted docx is not a zip, so it fails and gets skipped.
    """
    lines: list[str] = []
    with zipfile.ZipFile(path) as z:  # encrypted docx is not a zip -> BadZipFile -> skipped
        names = [n for n in z.namelist() if re.fullmatch(
            r"word/(document|footnotes|endnotes|header\d*|footer\d*)\.xml", n)]
        names.sort(key=lambda n: (n != "word/document.xml", n))
        for name in names:
            _docx_walk(ET.fromstring(z.read(name)), lines)
    return [{"text": t, "loc": f"lines {a}-{b}"} for a, b, t in chunk_lines(lines)]


def parse_pdf(path: Path, describe_image=None) -> list[dict]:
    """
    	Uses pypdf. An encrypted PDF raises an error, even if a blank password would open it. It extracts text page by page, with location p.N. If a PDF has almost no text (a scan),
        it renders up to 8 pages with pypdfium2 and asks the vision model to transcribe them.
    """
    from pypdf import PdfReader

    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        # Even with a blank user password: nothing inside an encrypted file is a source.
        raise PermissionError("encrypted pdf")
    out, total = [], 0
    n_pages = len(reader.pages)
    for pno, page in enumerate(reader.pages, 1):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        total += len(txt.strip())
        for a, b, t in chunk_lines(txt.splitlines()):
            out.append({"text": t, "loc": f"p.{pno}" + (f" lines {a}-{b}" if b - a > 0 and len(txt) > MAX_CHARS else "")})
    if total < 30 * max(1, n_pages) and describe_image is not None:  # scanned PDF -> vision
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(path))
            for pno in range(min(len(pdf), 8)):
                img = pdf[pno].render(scale=2).to_pil().convert("RGB")
                t = describe_image(img, f"{path.name} p.{pno + 1}")
                if t:
                    out.append({"text": t, "loc": f"p.{pno + 1} (scan)"})
        except Exception:
            pass
    return out


def load_image(path: Path):
    """
    Opens an image with PIL, applies its rotation metadata, and converts it to RGB.
    """
    from PIL import Image, ImageOps

    img = Image.open(path)
    img.load()
    img = ImageOps.exif_transpose(img).convert("RGB")
    return img


def parse_image(path: Path, describe_image=None) -> list[dict]:
    """
    Calls the vision model to transcribe the image into one chunk. If no transcription is available,
    the chunk just says IMAGE FILE (not transcribed), so the file is still indexed by its name.
    """
    text = ""
    if describe_image is not None:
        text = describe_image(load_image(path), path.name) or ""
    body = ("IMAGE FILE. Text and content visible in the image:\n" + text) if text else "IMAGE FILE (not transcribed)."
    return [{"text": body, "loc": "image"}]


# --------------------------------------------------------------------------- dispatch

def kind_of(path: Path) -> str | None:
    """
    Maps the file extension to a type (pdf, docx, xlsx, table, text, image),
    or None for unknown types, which get skipped.
    """
    ext = path.suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext == ".xlsx":
        return "xlsx"
    if ext in TABLE_EXT:
        return "table"
    if ext in TEXT_EXT:
        return "text"
    if ext in IMAGE_EXT:
        return "image"
    return None  # unknown type: skipped


def parse_file(path: Path, kind: str, describe_image=None) -> list[dict]:
    """
    The dispatcher. It calls the right parser for the type.
    """
    if kind == "pdf":
        return parse_pdf(path, describe_image)
    if kind == "docx":
        return parse_docx(path)
    if kind == "xlsx":
        return parse_xlsx(path)
    if kind == "table":
        return parse_csv(path)
    if kind == "text":
        return parse_text(path)
    if kind == "image":
        return parse_image(path, describe_image)
    raise ValueError(kind)


_STATUS_NAME = re.compile(r"(withdrawn|superseded|obsolete)", re.I)
_STATUS_BANNER = re.compile(
    r"(?im)^\W*(?:document\s+)?(?:status\s*[:\-]\s*)?(?:withdrawn|superseded|obsolete)\b"
    r"|\b(?:this|the)\s+(?:document|datasheet|revision|specification)\s+(?:is|has\s+been|was)\s+"
    r"(?:now\s+)?(?:withdrawn|superseded|obsolete)\b"
)


def file_status(rel: str, first_text: str) -> str:
    """
    Returns "withdrawn" if the filename contains withdrawn, superseded or obsolete, or if the first 600 characters
    carry a banner such as "THIS DATASHEET IS WITHDRAWN". Otherwise "current". It does not search the whole body,
    so a current file that says "this supersedes the withdrawn r1" isn't flagged by mistake.
    """
    if _STATUS_NAME.search(rel) or _STATUS_BANNER.search(first_text[:600]):
        return "withdrawn"
    return "current"
